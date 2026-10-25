# -*- coding: utf-8 -*-
"""
Generator FSD (Functional Specification Document)
Sistem: Cafe Monitoring & Enterprise ERP System (KopiERP Pro v2.4.8)
Format: Microsoft Word (.docx)
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    if top:
        borders_xml += f'<w:top w:val="{top.get("val", "single")}" w:sz="{top.get("sz", "4")}" w:space="0" w:color="{top.get("color", "CBD5E1")}"/>'
    if bottom:
        borders_xml += f'<w:bottom w:val="{bottom.get("val", "single")}" w:sz="{bottom.get("sz", "4")}" w:space="0" w:color="{bottom.get("color", "CBD5E1")}"/>'
    if left:
        borders_xml += f'<w:left w:val="{left.get("val", "single")}" w:sz="{left.get("sz", "4")}" w:space="0" w:color="{left.get("color", "CBD5E1")}"/>'
    if right:
        borders_xml += f'<w:right w:val="{right.get("val", "single")}" w:sz="{right.get("sz", "4")}" w:space="0" w:color="{right.get("color", "CBD5E1")}"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=6)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(17)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # #0F172A
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=10, space_after=4)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 64, 175) # #1E40AF
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=7, space_after=3)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85) # #334155
    return p

def add_body_p(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, title, text, alert_type="info"):
    colors = {
        "info": {"bg": "F0F9FF", "border": "0284C7", "title_rgb": RGBColor(2, 132, 199)},
        "warning": {"bg": "FFFBEB", "border": "D97706", "title_rgb": RGBColor(217, 119, 6)},
        "success": {"bg": "ECFDF5", "border": "059669", "title_rgb": RGBColor(5, 150, 105)},
        "danger": {"bg": "FEF2F2", "border": "DC2626", "title_rgb": RGBColor(220, 38, 38)}
    }
    cfg = colors.get(alert_type, colors["info"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_borders(cell, left={"val": "single", "sz": "24", "color": cfg["border"]},
                           top={"val": "none"}, bottom={"val": "none"}, right={"val": "none"})
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = cfg["title_rgb"]
    
    r_text = p.add_run(text)
    r_text.font.name = 'Segoe UI'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    format_paragraph(p_after, space_before=0, space_after=4)

def add_custom_table(doc, headers, data, col_widths=None, align_cols=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=0, space_after=0)
        run = p.add_run(title)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        if align_cols and i < len(align_cols):
            p.alignment = align_cols[i]
    
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_hex = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = row_cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            set_cell_borders(cell, top={"val": "single", "sz": "4", "color": "E2E8F0"},
                                   bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                                   left={"val": "single", "sz": "4", "color": "E2E8F0"},
                                   right={"val": "single", "sz": "4", "color": "E2E8F0"})
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=0, space_after=0)
            run = p.add_run(str(val))
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)
            if align_cols and col_idx < len(align_cols):
                p.alignment = align_cols[col_idx]
                
    p_after = doc.add_paragraph()
    format_paragraph(p_after, space_before=0, space_after=4)
    return table

def add_mockup_box(doc, screen_title, mockup_ascii_art, controls_info):
    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=6, space_after=2)
    run = p_title.add_run(f"🖥️ Mockup Wireframe Antarmuka: {screen_title}")
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Outer box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A") # Dark console / wireframe look
    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
    set_cell_borders(cell, top={"val": "single", "sz": "12", "color": "334155"},
                           bottom={"val": "single", "sz": "12", "color": "334155"},
                           left={"val": "single", "sz": "12", "color": "334155"},
                           right={"val": "single", "sz": "12", "color": "334155"})
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=0, line_spacing=1.0)
    run_ascii = p.add_run(mockup_ascii_art)
    run_ascii.font.name = 'Consolas'
    run_ascii.font.size = Pt(8)
    run_ascii.font.color.rgb = RGBColor(226, 232, 240)
    
    # Table of controls
    if controls_info:
        doc.add_paragraph()
        add_custom_table(
            doc,
            headers=["Komponen / Elemen UI", "Tipe Kontrol", "Aksi / Interaksi Pengguna", "Aturan & Efek Sistem"],
            data=controls_info,
            col_widths=[Inches(1.6), Inches(1.2), Inches(1.8), Inches(1.9)]
        )

def add_flowchart(doc, flow_name, ascii_chart, stages_data):
    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=6, space_after=2)
    run = p_title.add_run(f"📊 Diagram Alur Sistem (Flowchart): {flow_name}")
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 64, 175)
    
    # Diagram box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_borders(cell, top={"val": "single", "sz": "8", "color": "CBD5E1"},
                           bottom={"val": "single", "sz": "8", "color": "CBD5E1"},
                           left={"val": "single", "sz": "18", "color": "2563EB"},
                           right={"val": "single", "sz": "8", "color": "CBD5E1"})
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=0, line_spacing=1.05)
    run_chart = p.add_run(ascii_chart)
    run_chart.font.name = 'Consolas'
    run_chart.font.size = Pt(8)
    run_chart.font.color.rgb = RGBColor(15, 23, 42)
    
    if stages_data:
        doc.add_paragraph()
        add_custom_table(
            doc,
            headers=["Tahap / Langkah", "Aktor", "Input / Kondisi", "Proses & Mutasi Database", "Output / Respons"],
            data=stages_data,
            col_widths=[Inches(1.2), Inches(1.1), Inches(1.3), Inches(1.6), Inches(1.3)]
        )

print("Helper definitions loaded successfully.")
